"""Verify generated artifacts in an isolated SQLite database, never the user's data."""
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

from openpyxl import load_workbook
from app.db.connection import make_engine, session_factory
from app.db.models import Base, Business
from app.features.credit.repository import credit_inputs
from app.features.credit.router import build_credit_snapshot
from app.integrations.excel.import_service import import_validated_snapshot
from app.integrations.excel.workbook_parser import parse_workbook
from app.integrations.excel.validators import validate_workbook

folder = ROOT / "outputs/credit-hub-20261022"
template = (folder / "credit-import-template.xlsx").read_bytes()
assert validate_workbook(parse_workbook(template), "YOUR_BUSINESS_ID") == []
sample = folder / "credit-import-synthetic-sample.xlsx"
parsed = parse_workbook(sample.read_bytes())
assert validate_workbook(parsed, "grocery") == []

# Round-trip preservation: unrelated sheets and original loan facts stay intact.
original = load_workbook(ROOT / "database/seed/grocery.xlsx")
updated = load_workbook(sample)
assert original.sheetnames == updated.sheetnames
for name in original.sheetnames:
    if name in {"Loans", "ReadMe"}:
        continue
    def populated_rows(sheet):
        return [tuple(row) for row in sheet.values if any(value is not None for value in row)]
    assert populated_rows(original[name]) == populated_rows(updated[name]), name
for column in (1, 2, 3, 5, 6):
    assert original["Loans"].cell(2, column).value == updated["Loans"].cell(2, column).value
original.close()
updated.close()

# In-memory DB makes repeat runs safe and excludes any production app data.
engine = make_engine("sqlite://")
Base.metadata.create_all(engine)
sessions = session_factory(engine)
with sessions.begin() as session:
    session.add(Business(id="grocery", name="Synthetic import check", industry="Retail", city="Yangon"))
result = import_validated_snapshot(sessions, "grocery", "artifact-check", sample.read_bytes())
with sessions() as session:
    snapshot = build_credit_snapshot(credit_inputs(session, "grocery"), "grocery")
loan = snapshot["loans"][0]
assert loan["bank_id"] == "kbz"
assert loan["loan_product_id"] == "kbz-term"
assert loan["annual_rate_pct"] == 12
assert loan["schedule_status"] == "available"
assert len(loan["schedule"]) == 24
assert sum(row.principal_due_mmk for row in loan["schedule"]) == loan["principal_mmk"]
assert loan["schedule"][0].due_date.isoformat() == "2026-11-01"
assert loan["schedule"][-1].projected_principal_mmk == 0
assert import_validated_snapshot(sessions, "grocery", "artifact-check", sample.read_bytes())["status"] == "up_to_date"
print("Template valid; sample imported; loan terms and 24 installments verified; re-import adds no duplicates; unrelated sheets preserved.")
engine.dispose()
