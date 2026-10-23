"""Read-only artifact audit and isolated import test."""
from pathlib import Path
import hashlib
import json
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))
from openpyxl import load_workbook
from app.integrations.excel.workbook_parser import parse_workbook
from app.integrations.excel.validators import validate_workbook
from app.integrations.excel.import_service import import_validated_snapshot
from app.db.connection import make_engine, session_factory
from app.db.models import Base, Business
from app.features.credit.repository import credit_inputs
from app.features.credit.router import build_credit_snapshot

folder = ROOT / "outputs/mi-thar-su-credit-20261022"
audit = json.loads((folder / "source-check.json").read_text())
assert hashlib.sha256(Path(audit["sourcePath"]).read_bytes()).hexdigest() == audit["sourceSHA256"]
source = load_workbook(audit["sourcePath"], read_only=True)
output = folder / "Mi-Thar-Su-KBZ-demo-import.xlsx"
target = load_workbook(output, read_only=True)
business_id = target["Settings"]["A2"].value
assert business_id == "0c009beb-8c8b-4445-a898-da7ab6661980"
assert target["Settings"]["B2"].value == "Mi Thar Su"
counts = 0
for name in source.sheetnames:
    if name == "ReadMe":
        continue
    source_rows = [row for row in source[name].values if any(value is not None for value in row)]
    target_rows = [row for row in target[name].values if any(value is not None for value in row)]
    headers = source_rows[0]
    expected = source_rows[1:]
    if name == "Loans":
        expected = [row for row in expected if row[0] == "LOAN-001"]
    elif name == "LoanPayments":
        expected = [row for row in expected if row[3] == "LOAN-001"]
    assert len(expected) == len(target_rows) - 1, name
    for src, dst in zip(expected, target_rows[1:]):
        for column, header in enumerate(headers):
            if header == "business_id":
                assert dst[column] == business_id
            elif name == "Settings" and header == "business_name":
                assert dst[column] == "Mi Thar Su"
            else:
                assert src[column] == dst[column], (name, header, src[column], dst[column])
        counts += 1
source.close()
target.close()

content = output.read_bytes()
assert validate_workbook(parse_workbook(content), business_id) == []
engine = make_engine("sqlite://")
Base.metadata.create_all(engine)
sessions = session_factory(engine)
with sessions.begin() as session:
    session.add(Business(id=business_id, name="Mi Thar Su", industry="Wholesale", city="Yangon"))
import_validated_snapshot(sessions, business_id, "artifact-verification", content)
with sessions() as session:
    loan = build_credit_snapshot(credit_inputs(session, business_id), business_id)["loans"][0]
assert loan["loan_id"] == "LOAN-001"
assert loan["schedule_status"] == "available"
assert len(loan["actual_payments"]) == 2
assert loan["actual_principal_paid_mmk"] == 800000
assert loan["actual_interest_paid_mmk"] == 355000
assert loan["outstanding_principal_mmk"] == 11200000
assert len(loan["schedule"]) == 24
assert loan["schedule"][0].due_date.isoformat() == "2026-11-03"
assert sum(row.principal_due_mmk for row in loan["schedule"]) == 12000000
assert loan["schedule"][-1].projected_principal_mmk == 0
assert import_validated_snapshot(sessions, business_id, "artifact-verification", content)["status"] == "up_to_date"
print(f"Verified {counts} records. Source unchanged. Copied facts match. Isolated import and 24-installment schedule pass. Re-import is idempotent.")
engine.dispose()
