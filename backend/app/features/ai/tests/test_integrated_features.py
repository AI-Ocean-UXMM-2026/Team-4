"""Exercise all four mounted features against the same uploaded workbook."""

from io import BytesIO

from openpyxl import load_workbook

from app.core.demo_data import is_synthetic_seed_workbook
from app.features.ai.context_builder import build_context
from app.features.ai.snapshot_provider import SnapshotProvider
from app.integrations.excel.tests.conftest import client, upload, workbook_bytes


def test_pinned_local_demo_requires_exact_business_and_import_hash(monkeypatch):
    digest = "a" * 64
    monkeypatch.setenv("AI_SYNTHETIC_DEMO_BUSINESS_ID", "demo-business")
    monkeypatch.setenv("AI_SYNTHETIC_DEMO_IMPORT_SHA256", digest)
    assert is_synthetic_seed_workbook("demo-business", digest)
    assert not is_synthetic_seed_workbook("another-business", digest)
    assert not is_synthetic_seed_workbook("demo-business", "b" * 64)
    monkeypatch.setenv("AI_SYNTHETIC_DEMO_IMPORT_SHA256", "invalid")
    assert not is_synthetic_seed_workbook("demo-business", digest)


def test_feature_routes_and_business_authorization(client):
    modules = client.get("/health").json()["modules"]
    assert all(modules[name] for name in ("financial", "working_capital", "credit", "ai"))
    empty_credit = client.get("/api/v1/credit/summary?business_id=tea-shop")
    assert empty_credit.status_code == 200
    assert empty_credit.json()["readiness_progress"]["prepared_categories"] == 0
    scenario = {"scenario_type": "receivable_collection",
                "assumptions": {"outstanding_mmk": "1000", "collection_pct": "50"}}
    assert client.post("/api/v1/working-capital/simulate", json=scenario).status_code == 400
    assert client.post("/api/v1/working-capital/simulate?business_id=private", json=scenario).status_code == 403
    response = client.post("/api/v1/working-capital/simulate?business_id=tea-shop", json=scenario)
    assert response.status_code == 200, response.text
    assert response.json()["records_changed"] is False


def test_uploaded_seed_connects_snapshots_and_proves_synthetic_origin(client, workbook_bytes):
    assert upload(client, workbook_bytes).status_code == 200
    version = client.get("/api/v1/data/snapshot?business_id=tea-shop").json()["source_version"]
    working = client.get("/api/v1/working-capital/summary?business_id=tea-shop")
    credit = client.get("/api/v1/credit/summary?business_id=tea-shop")
    assert working.status_code == credit.status_code == 200
    assert working.json()["source_version"] == credit.json()["source_version"] == version
    provider = SnapshotProvider(client.app.state.sessions)
    bundle = provider.load("tea-shop")
    assert bundle.financial and bundle.working_capital and bundle.credit
    assert build_context(bundle)["synthetic"] is True

    book = load_workbook(BytesIO(workbook_bytes))
    sheet = book["Expenses"]
    amount_column = [cell.value for cell in sheet[1]].index("amount_mmk") + 1
    sheet.cell(2, amount_column).value = float(sheet.cell(2, amount_column).value) + 1
    changed = BytesIO()
    book.save(changed)
    book.close()
    assert upload(client, changed.getvalue()).status_code == 200
    assert build_context(provider.load("tea-shop"))["synthetic"] is False
