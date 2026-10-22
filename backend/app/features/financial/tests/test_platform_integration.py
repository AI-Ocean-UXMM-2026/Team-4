"""Exercise actual P5 imports, SQL projections and business authorization."""
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from openpyxl import load_workbook
from sqlalchemy import select, update

from app.db.connection import make_engine, session_factory
from app.db.models import Base, Business, ImportedRow, PROJECTIONS, WorkbookSource
from app.integrations.excel.tests.conftest import ROOT, client, edit_workbook, workbook_bytes, upload
from app.integrations.excel.import_service import import_validated_snapshot
from app.integrations.excel.router import snapshot as data_snapshot
from app.features.financial.platform import calculate_platform_snapshot
from app.features.financial.platform_repository import Ledger, SqlFinancialRepository
from app.features.financial.schemas import Period
from datetime import date


URL = "/api/v1/financial/summary?business_id=tea-shop&start=2026-10-01&end=2026-10-31"


def test_real_import_summary_and_idempotency(client, workbook_bytes):
    first = upload(client, workbook_bytes)
    assert first.status_code == 200
    response = client.get(URL)
    assert response.status_code == 200, response.text
    summary = response.json()
    assert summary["synthetic"] is True
    source = client.get("/api/v1/data/snapshot?business_id=tea-shop").json()
    assert source["synthetic"] is True
    sales = source["tables"]["Sales"]
    expected = sum((Decimal(x["quantity"]) * Decimal(x["unit_price_mmk"]) for x in sales), Decimal(0))
    assert Decimal(summary["actual"]["revenue_mmk"]) == expected
    assert summary["actual"]["cogs_mmk"] is None
    assert summary["actual"]["profit_before_tax_mmk"] is None
    assert summary["actual"]["cash_balance_mmk"] is None
    assert summary["source_version"] == source["source_version"]
    assert summary["as_of"] and summary["import_batch_id"]
    trends = client.get("/api/v1/financial/trends?business_id=tea-shop&end=2026-10-20&months=3")
    assert trends.status_code == 200, trends.text
    august, september, october = trends.json()["points"]
    # Empty imported months are identified as no records, not as recorded 0 MMK.
    assert august["revenue_record_count"] == august["operating_expense_record_count"] == 0
    assert september["revenue_record_count"] == september["operating_expense_record_count"] == 0
    assert october["revenue_record_count"] == len(sales) == 3
    assert october["operating_expense_record_count"] == 2
    assert october["period"] == {"start": "2026-10-01", "end": "2026-10-20"}
    assert upload(client, workbook_bytes).status_code == 200
    assert client.get(URL).json() == summary


def test_only_exact_seed_workbook_is_marked_synthetic(tmp_path):
    engine = make_engine(f"sqlite:///{tmp_path / 'synthetic.db'}")
    Base.metadata.create_all(engine)
    sessions = session_factory(engine)
    with sessions.begin() as session:
        session.add(Business(id="tea-shop", name="Tea Shop", industry="Food", city="Yangon"))
    content = (ROOT / "database/seed/tea-shop.xlsx").read_bytes()
    import_validated_snapshot(sessions, "tea-shop", "owner", content)
    repository = SqlFinancialRepository(sessions)
    assert repository.load_ledger("tea-shop").synthetic is True
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(sessions=sessions)))
    assert data_snapshot(request, "tea-shop", "owner")["synthetic"] is True

    changed = edit_workbook(content, lambda w: setattr(w.properties, "title", "Edited copy"))
    import_validated_snapshot(sessions, "tea-shop", "owner", changed)
    assert repository.load_ledger("tea-shop").synthetic is False
    assert data_snapshot(request, "tea-shop", "owner")["synthetic"] is False
    engine.dispose()


def test_wholesale_october_workbook_import_and_financial_availability(client):
    path = Path(__file__).parent / "fixtures" / "BIZFLOW_Wholesale_Sample_October_2026.xlsx"
    workbook = load_workbook(BytesIO(path.read_bytes()))
    # Bind the synthetic sample to this test user's authorized business.
    for sheet in workbook.worksheets:
        headers = [cell.value for cell in sheet[1]]
        if "business_id" in headers:
            column = headers.index("business_id") + 1
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, column).value = "tea-shop"
    content = BytesIO()
    workbook.save(content)
    workbook.close()

    imported = upload(client, content.getvalue())
    assert imported.status_code == 200, imported.text
    assert imported.json()["status"] == "success"

    response = client.get(
        "/api/v1/financial/summary?business_id=tea-shop&start=2026-10-01&end=2026-10-21"
    )
    assert response.status_code == 200, response.text
    summary = response.json()
    assert summary["synthetic"] is False
    actual = summary["actual"]
    assert Decimal(actual["revenue_mmk"]) == Decimal("13061400")
    assert Decimal(actual["operating_expenses_mmk"]) == Decimal("2657000")
    assert Decimal(actual["interest_expenses_mmk"]) == Decimal("180000")
    assert Decimal(actual["receivables_mmk"]) == Decimal("2912900")
    assert Decimal(actual["payables_mmk"]) == Decimal("5051200")

    # The imported v1 sheets do not state expense settlement or cost at sale.
    # Keep the opening cash, product costs and purchase costs without presenting
    # estimates as actual cash or historical COGS/profit.
    assert actual["cash_balance_mmk"] is None
    assert actual["cogs_mmk"] is None
    assert actual["operating_profit_mmk"] is None
    warnings = summary["data_quality"]["warnings"]
    assert any("Expense cash movements are incomplete" in item for item in warnings)
    assert any("Historical cost is missing" in item for item in warnings)

    snapshot = client.get("/api/v1/data/snapshot?business_id=tea-shop").json()
    tables = snapshot["tables"]
    assert tables["Settings"][0]["opening_cash_mmk"] == "8500000.00"
    assert len(tables["Products"]) == 12
    assert len(tables["Sales"]) == 48
    assert len(tables["Purchases"]) == 28
    assert "Rent" in {row["category"] for row in tables["Expenses"]}
    assert {row["loan_type"] for row in tables["Loans"]} == {"Business expansion", "Working capital"}
    assert "partial" not in {row["payment_type"] for row in tables["Sales"]}

    credit_loans = client.get("/api/v1/credit/loans?business_id=tea-shop")
    assert credit_loans.status_code == 200, credit_loans.text
    assert all(item["schedule_status"] == "unsupported_loan_model" for item in credit_loans.json()["loans"])
    assert all(item["schedule"] == [] for item in credit_loans.json()["loans"])


def test_enhanced_wholesale_workbook_imports_costs_and_dated_expense_cash(client):
    path = Path(__file__).parent / "fixtures" / "BIZFLOW_Wholesale_Sample_October_2026_Enhanced.xlsx"
    workbook = load_workbook(BytesIO(path.read_bytes()))
    for sheet in workbook.worksheets:
        headers = [cell.value for cell in sheet[1]]
        if "business_id" in headers:
            column = headers.index("business_id") + 1
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, column).value = "tea-shop"
    content = BytesIO()
    workbook.save(content)
    workbook.close()

    imported = upload(client, content.getvalue())
    assert imported.status_code == 200, imported.text
    summary_response = client.get(
        "/api/v1/financial/summary?business_id=tea-shop&start=2026-10-01&end=2026-10-21"
    )
    assert summary_response.status_code == 200, summary_response.text
    summary = summary_response.json()
    actual = summary["actual"]
    assert Decimal(actual["revenue_mmk"]) == Decimal("13061400.00")
    assert Decimal(actual["cogs_mmk"]) == Decimal("11079500.00")
    assert Decimal(actual["operating_expenses_mmk"]) == Decimal("2657000.00")
    assert Decimal(actual["operating_profit_mmk"]) == Decimal("-675100.00")
    assert Decimal(actual["cash_balance_mmk"]) == Decimal("12919300.00")

    warnings = summary["data_quality"]["warnings"]
    assert any("quantity × historical unit cost" in warning for warning in warnings)
    assert any("less recorded payments through 2026-10-21" in warning and "paid date" in warning for warning in warnings)

    snapshot = client.get("/api/v1/data/snapshot?business_id=tea-shop").json()["tables"]
    assert snapshot["Sales"][0]["unit_cost_at_sale_mmk"] is not None
    assert snapshot["Sales"][0]["cogs_mmk"] is not None
    partial = next(row for row in snapshot["Expenses"] if row["payment_status"] == "partial")
    assert partial["paid_date"] is not None
    assert Decimal(partial["amount_paid_mmk"]) < Decimal(partial["amount_mmk"])


def test_same_hash_reimport_backfills_fields_added_after_legacy_import(client):
    path = Path(__file__).parent / "fixtures" / "BIZFLOW_Wholesale_Sample_October_2026_Enhanced.xlsx"
    workbook = load_workbook(BytesIO(path.read_bytes()))
    for sheet in workbook.worksheets:
        headers = [cell.value for cell in sheet[1]]
        if "business_id" in headers:
            column = headers.index("business_id") + 1
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, column).value = "tea-shop"
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    uploaded = content.getvalue()
    assert upload(client, uploaded).status_code == 200

    # Recreate rows imported by the older parser while leaving the file hash unchanged.
    with client.app.state.sessions.begin() as session:
        source = session.scalar(select(WorkbookSource).where(WorkbookSource.business_id == "tea-shop"))
        rows = session.scalars(select(ImportedRow).where(ImportedRow.business_id == "tea-shop", ImportedRow.source_id == source.id)).all()
        for row in rows:
            optional = {"Sales": ("unit_cost_at_sale_mmk", "cogs_mmk"), "Expenses": ("payment_status", "paid_date", "amount_paid_mmk")}.get(row.sheet_name)
            if optional is None:
                continue
            payload = dict(row.payload)
            for field in optional:
                payload.pop(field, None)
            row.payload = payload
            table = PROJECTIONS[row.sheet_name]
            session.execute(update(table).where(table.c.source_id == source.id, table.c["sale_id" if row.sheet_name == "Sales" else "expense_id"] == row.row_id).values(**{field: None for field in optional}))

    refreshed = upload(client, uploaded)
    assert refreshed.status_code == 200, refreshed.text
    assert refreshed.json()["status"] == "success"
    assert refreshed.json()["row_counts"]["updated"] == 66

    summary = client.get("/api/v1/financial/summary?business_id=tea-shop&start=2026-10-01&end=2026-10-21").json()
    assert Decimal(summary["actual"]["cash_balance_mmk"]) == Decimal("12919300.00")
    assert Decimal(summary["actual"]["cogs_mmk"]) == Decimal("11079500.00")
    assert Decimal(summary["actual"]["operating_profit_mmk"]) == Decimal("-675100.00")


def test_edit_expense_reimports_once(client, workbook_bytes):
    assert upload(client, workbook_bytes).status_code == 200
    before = client.get(URL).json()
    workbook = load_workbook(BytesIO(workbook_bytes))
    sheet = workbook["Expenses"]
    column = [cell.value for cell in sheet[1]].index("amount_mmk") + 1
    sheet.cell(2, column).value = float(Decimal(str(sheet.cell(2, column).value)) + Decimal("30000"))
    target = BytesIO()
    workbook.save(target)
    workbook.close()
    changed = upload(client, target.getvalue())
    assert changed.status_code == 200, changed.text
    assert changed.json()["row_counts"]["updated"] == 1
    after = client.get(URL).json()
    assert Decimal(after["actual"]["operating_expenses_mmk"]) - Decimal(before["actual"]["operating_expenses_mmk"]) == 30000
    assert after["source_version"] > before["source_version"]
    assert after["import_batch_id"] != before["import_batch_id"]


def test_authorization_and_no_import(client):
    assert client.get(URL).json()["data_quality"]["status"] == "missing_data"
    assert client.get(URL.replace("tea-shop", "private")).status_code == 403
    client.headers.pop("Authorization")
    assert client.get(URL).status_code == 401


def test_trends_and_undated_obligations(client, workbook_bytes):
    upload(client, workbook_bytes)
    trends = client.get("/api/v1/financial/trends?business_id=tea-shop&end=2026-10-31&months=3")
    assert trends.status_code == 200
    assert len(trends.json()["points"]) == 3
    obligations = client.get("/api/v1/financial/obligations?business_id=tea-shop&as_at=2026-10-31&before=2026-11-30")
    assert obligations.status_code == 200
    assert all(x["due_date"] is None for x in obligations.json()["upcoming_obligations"])


def test_partial_initial_settlement_decimal_cash_and_receivables():
    d = date(2026, 10, 1)
    tables = {name: [] for name in ["Settings", "Sales", "CustomerReceipts", "Purchases", "SupplierPayments", "Expenses", "Loans", "LoanPayments", "CashAdjustments"]}
    tables["Settings"] = [dict(business_id="A", opening_date=d, opening_cash_mmk="1000.00")]
    tables["Sales"] = [dict(business_id="A", sale_id="S", date=d, quantity="1.500", unit_price_mmk="100.50", cash_received_mmk="50.25")]
    tables["CustomerReceipts"] = [dict(business_id="A", receipt_id="R", sale_id="S", date=d, amount_mmk="20.10")]
    tables["Purchases"] = [dict(business_id="A", purchase_id="P", date=d, quantity="1.000", unit_cost_mmk="30.25", cash_paid_mmk="10.25")]
    ledger = Ledger("A", 1, tables, None, "batch")
    snap = calculate_platform_snapshot(ledger, Period(start=d, end=d))
    assert snap.actual.revenue_mmk == Decimal("150.75")
    assert snap.actual.receivables_mmk == Decimal("80.40")
    assert snap.actual.payables_mmk == Decimal("20.00")
    assert snap.actual.cash_balance_mmk == Decimal("1060.10")
    assert snap.upcoming_obligations[0].due_date is None


def test_cogs_falls_back_to_historical_unit_cost_once():
    d = date(2026, 10, 1)
    tables = {name: [] for name in ["Settings", "Sales", "CustomerReceipts", "Purchases", "SupplierPayments", "Expenses", "Loans", "LoanPayments", "CashAdjustments"]}
    tables["Settings"] = [dict(business_id="A", opening_date=d, opening_cash_mmk="1000.00")]
    tables["Sales"] = [dict(business_id="A", sale_id="S", date=d, quantity="2.000", unit_price_mmk="100.00", cash_received_mmk="200.00", unit_cost_at_sale_mmk="30.00", cogs_mmk=None)]
    snap = calculate_platform_snapshot(Ledger("A", 1, tables, None, "batch"), Period(start=d, end=d))
    assert snap.actual.cogs_mmk == Decimal("60.00")
    assert snap.actual.operating_profit_mmk == Decimal("140.00")


def test_future_unpaid_expense_does_not_invalidate_earlier_cash_balance():
    from datetime import date
    d = date(2026, 10, 1)
    tables = {name: [] for name in ["Settings", "Sales", "CustomerReceipts", "Purchases", "SupplierPayments", "Expenses", "Loans", "LoanPayments", "CashAdjustments"]}
    tables["Settings"] = [dict(business_id="A", opening_date=date(2026, 1, 1), opening_cash_mmk="1000.00")]
    tables["Expenses"] = [
        dict(business_id="A", expense_id="past", date=date(2026, 9, 1), amount_mmk="100.00", amount_paid_mmk="100.00", paid_date=date(2026, 9, 2)),
        dict(business_id="A", expense_id="future", date=date(2026, 11, 1), amount_mmk="200.00", amount_paid_mmk=None, paid_date=None),
    ]
    snap = calculate_platform_snapshot(Ledger("A", 1, tables, None, "batch"), Period(start=d, end=d))
    assert snap.actual.cash_balance_mmk == Decimal("900.00")
    assert not any("Expense cash movements are incomplete" in warning for warning in snap.data_quality.warnings)
