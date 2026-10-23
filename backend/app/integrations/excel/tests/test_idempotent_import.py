from concurrent.futures import ThreadPoolExecutor

import pytest
from sqlalchemy import func, select

from app.core.errors import APIError
from app.db.connection import make_engine, session_factory
from app.db.models import Base, Business, ImportBatch, ImportedRow, PROJECTIONS, RowRevision
from app.integrations.excel.import_service import import_validated_snapshot
from conftest import ROOT, edit_workbook, upload


def snapshot(client):
    return client.get("/api/v1/data/snapshot?business_id=tea-shop").json()


def test_repeat_and_single_expense_edit(client, workbook_bytes):
    initial = upload(client, workbook_bytes)
    assert initial.status_code == 200, initial.text
    before = snapshot(client)
    assert upload(client, workbook_bytes).json()["status"] == "up_to_date"
    assert snapshot(client) == before
    changed = edit_workbook(workbook_bytes, lambda w: setattr(w["Expenses"]["E2"], "value", 80000))
    result = upload(client, changed)
    assert result.status_code == 200, result.text
    assert result.json()["row_counts"] == {"inserted": 0, "updated": 1, "unchanged": 10}
    after = snapshot(client)
    assert after["source_version"] == before["source_version"] + 1
    assert len(after["tables"]["Expenses"]) == 2
    assert after["tables"]["Expenses"][0]["amount_mmk"] == "80000.00"
    with client.app.state.sessions() as session:
        assert session.scalar(select(func.count()).select_from(RowRevision)) == 1
        assert session.scalar(select(RowRevision)).payload["amount_mmk"] == "50000.00"


def test_duplicate_fails_atomically_and_is_audited(client, workbook_bytes):
    upload(client, workbook_bytes)
    before = snapshot(client)
    def change(w):
        w["Sales"]["A3"] = "S001"
        w["Expenses"]["E2"] = 80000
    result = upload(client, edit_workbook(workbook_bytes, change))
    assert result.status_code == 400
    assert snapshot(client) == before
    status = client.get("/api/v1/data/sync-status?business_id=tea-shop").json()
    assert status["imports"][0]["status"] == "failed"
    assert status["last_import_at"]


@pytest.mark.parametrize("sheet,cell,value", [("CustomerReceipts", "E2", 90000), ("SupplierPayments", "E2", 110000), ("Sales", "E2", 9999), ("Sales", "H2", 1), ("Sales", "D2", "foreign-id")])
def test_invalid_financial_state_rejected(client, workbook_bytes, sheet, cell, value):
    upload(client, workbook_bytes)
    before = snapshot(client)
    data = edit_workbook(workbook_bytes, lambda w: setattr(w[sheet][cell], "value", value))
    assert upload(client, data).status_code == 400
    assert snapshot(client) == before


def test_no_cross_business_access(client, workbook_bytes):
    assert client.get("/api/v1/data/snapshot?business_id=private").status_code == 403
    assert client.post("/api/v1/data/upload?business_id=private", files={"file": ("book.xlsx", workbook_bytes)}).status_code == 403
    assert snapshot(client)["source_version"] == 0


def test_missing_rows_retained_and_checked(client, workbook_bytes):
    upload(client, workbook_bytes)
    data = edit_workbook(workbook_bytes, lambda w: w["Expenses"].delete_rows(3))
    result = upload(client, data)
    assert result.status_code == 200
    assert "Missing row E002 retained" in str(result.json()["warnings"])
    assert len(snapshot(client)["tables"]["Expenses"]) == 2


def test_replacement_reactivates_reappearing_row(tmp_path):
    engine = make_engine(f"sqlite:///{tmp_path / 'replacement.db'}")
    Base.metadata.create_all(engine)
    sessions = session_factory(engine)
    with sessions.begin() as session:
        session.add(Business(id="tea-shop", name="Tea Shop", industry="Food", city="Yangon"))
    content = (ROOT / "database/seed/tea-shop.xlsx").read_bytes()
    shortened = edit_workbook(content, lambda w: w["Expenses"].delete_rows(3))
    import_validated_snapshot(sessions, "tea-shop", "owner", content)
    merged = import_validated_snapshot(sessions, "tea-shop", "owner", shortened)
    assert "Missing row E002 retained" in str(merged["warnings"])
    result = import_validated_snapshot(sessions, "tea-shop", "owner", shortened, replace_omitted=True)
    assert result["row_counts"]["removed"] == 1
    with sessions() as session:
        omitted = session.scalar(select(ImportedRow).where(ImportedRow.row_id == "E002"))
        assert omitted.active_status == "removed"
        assert session.scalar(select(func.count()).select_from(PROJECTIONS["Expenses"])) == 1
        assert session.scalar(select(func.count()).select_from(RowRevision)) == 1
    restored = import_validated_snapshot(sessions, "tea-shop", "owner", content)
    assert restored["row_counts"]["updated"] == 1
    with sessions() as session:
        row = session.scalar(select(ImportedRow).where(ImportedRow.row_id == "E002"))
        assert row.active_status == "active"
        assert session.scalar(select(func.count()).select_from(PROJECTIONS["Expenses"])) == 2
    engine.dispose()


def test_omitted_payments_still_prevent_overpayment(client, workbook_bytes):
    upload(client, workbook_bytes)
    def change(w):
        w["Sales"]["E3"] = 5  # total 25000, initial paid 20000, retained receipt is 30000
        w["CustomerReceipts"].delete_rows(2)
    assert upload(client, edit_workbook(workbook_bytes, change)).status_code == 400


def test_sql_failure_rolls_back_every_fact(client, workbook_bytes, monkeypatch):
    upload(client, workbook_bytes)
    before = snapshot(client)
    from app.integrations.excel import import_service
    real = import_service.typed_values
    def fail(sheet, payload):
        if sheet == "Expenses" and payload["expense_id"] == "E002":
            raise RuntimeError("simulated database projection failure")
        return real(sheet, payload)
    monkeypatch.setattr(import_service, "typed_values", fail)
    def change(w):
        w["Expenses"]["E2"] = 80000
        w["Expenses"]["E3"] = 99999
    with pytest.raises(RuntimeError):
        import_validated_snapshot(client.app.state.sessions, "tea-shop", "demo-user", edit_workbook(workbook_bytes, change))
    assert snapshot(client) == before


def test_concurrent_same_file_import_is_serialized(client, workbook_bytes):
    def run(_):
        return import_validated_snapshot(client.app.state.sessions, "tea-shop", "demo-user", workbook_bytes)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(run, range(2)))
    assert sorted(r["status"] for r in results) == ["success", "up_to_date"]
    assert len(snapshot(client)["tables"]["Expenses"]) == 2


def test_template_is_personalized_and_public_reset_removed(client, workbook_bytes):
    template = client.get("/api/v1/data/template?business_id=tea-shop")
    assert template.status_code == 200
    from openpyxl import load_workbook
    from io import BytesIO
    workbook = load_workbook(BytesIO(template.content), read_only=True)
    assert "due_date" in [cell.value for cell in workbook["Purchases"][1]]
    assert all("business_id" not in [cell.value for cell in sheet[1]] for sheet in workbook)
    workbook.close()
    from app.integrations.excel.workbook_parser import parse_workbook
    assert parse_workbook(template.content).sheets["Settings"][0].values["business_name"] == "Tea Shop"
    def complete_opening_values(workbook):
        sheet = workbook["Settings"]
        columns = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(2, columns["opening_date"], "2026-10-01")
        sheet.cell(2, columns["opening_cash_mmk"], 1000000)
    blank_result = upload(client, edit_workbook(template.content, complete_opening_values))
    assert blank_result.status_code == 200, blank_result.text
    assert snapshot(client)["tables"]["Settings"][0]["business_id"] == "tea-shop"
    changed = edit_workbook(workbook_bytes, lambda w: setattr(w["Expenses"]["E2"], "value", 80000))
    upload(client, changed)
    result = client.post("/api/v1/demo/reset?business_id=tea-shop")
    assert result.status_code == 404
    assert snapshot(client)["tables"]["Expenses"][0]["amount_mmk"] == "80000.00"


def test_auth_and_upload_limits(client):
    assert client.post("/api/v1/data/upload?business_id=tea-shop", files={"file": ("book.csv", b"a,b")}).status_code == 400
    client.app.state.settings.__dict__["max_upload_bytes"] = 2
    assert client.post("/api/v1/data/upload?business_id=tea-shop", files={"file": ("book.xlsx", b"123")}).status_code == 400
    client.headers.pop("Authorization")
    assert client.get("/api/v1/data/snapshot?business_id=tea-shop").status_code == 401
