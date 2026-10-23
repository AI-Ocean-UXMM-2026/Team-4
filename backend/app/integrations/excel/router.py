from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from copy import copy

from fastapi import APIRouter, Depends, File, Query, Request, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import select

from ...core.dependencies import authorize_business
from ...core.demo_data import is_synthetic_seed_workbook
from ...core.errors import APIError
from ...db.models import Business, ImportBatch, PROJECTIONS, WorkbookSource
from .graph_client import GraphClient
from .import_audit import batch_response, iso_timestamp
from .import_service import import_validated_snapshot, lock_business
from .purchase_due_dates import change_purchase_due_date
from .onedrive_auth import access_token, require_graph
from .sync_service import sync_now
from .schemas import SCHEMA
from .workbook_storage import get_batch_workbook

router = APIRouter(prefix="/api/v1/data", tags=["data sources"])
ROOT = Path(__file__).resolve().parents[4]


@router.post("/upload")
def upload(request: Request, business_id: str, file: UploadFile = File(...), replace_omitted: bool = Query(False), due_date_resolution: str | None = Query(None), user_id: str = Depends(authorize_business)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise APIError(400, "FILE_TYPE_INVALID", "Choose an .xlsx workbook.")
    content = file.file.read(request.app.state.settings.max_upload_bytes + 1)
    if len(content) > request.app.state.settings.max_upload_bytes:
        raise APIError(400, "FILE_TOO_LARGE", "Workbook exceeds the upload size limit.")
    return import_validated_snapshot(request.app.state.sessions, business_id, user_id, content, replace_omitted=replace_omitted, filename=file.filename, due_date_resolution=due_date_resolution)


class DueDateUpdate(BaseModel):
    due_date: date
    expected_due_date: date | None = None


@router.put("/purchases/{purchase_id}/due-date")
def update_purchase_due_date(request: Request, purchase_id: str, change: DueDateUpdate,
                             business_id: str, user_id: str = Depends(authorize_business)):
    return change_purchase_due_date(request.app.state.sessions, business_id, purchase_id,
                                    user_id, change.due_date, change.expected_due_date)


@router.post("/activate-history")
def activate_history(request: Request, business_id: str, batch_id: str = Query(...), user_id: str = Depends(authorize_business)):
    with request.app.state.sessions() as session:
        batch = session.get(ImportBatch, batch_id)
        if not batch or batch.business_id != business_id:
            raise APIError(404, "BATCH_NOT_FOUND", "Import batch not found.")
        if batch.status not in ("success", "up_to_date"):
            raise APIError(400, "BATCH_NOT_RESTORABLE", "Only successful uploads can be activated.")
        filename = (batch.row_counts.get("filename") if isinstance(batch.row_counts, dict) else None) or f"snapshot-{batch.id[:8]}.xlsx"
        content = get_batch_workbook(business_id, batch.id)
        if not content:
            raise APIError(404, "WORKBOOK_FILE_NOT_FOUND", "Historical workbook file is no longer on disk.")
    result = import_validated_snapshot(request.app.state.sessions, business_id, user_id, content, replace_omitted=True, filename=filename)
    return {"status": "activated", "batch_id": batch.id, "filename": filename, "result": result}


@router.get("/history/{batch_id}/download")
def download_history(request: Request, business_id: str, batch_id: str, _user_id: str = Depends(authorize_business)):
    with request.app.state.sessions() as session:
        batch = session.get(ImportBatch, batch_id)
        if not batch or batch.business_id != business_id:
            raise APIError(404, "BATCH_NOT_FOUND", "Import batch not found.")
        content = get_batch_workbook(business_id, batch.id)
        if not content:
            raise APIError(404, "WORKBOOK_FILE_NOT_FOUND", "Historical workbook file is not available.")
        filename = (batch.row_counts.get("filename") if isinstance(batch.row_counts, dict) else None) or f"workbook-{batch.id[:8]}.xlsx"
        return Response(content, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.post("/sync-now")
def sync(request: Request, business_id: str, user_id: str = Depends(authorize_business)):
    require_graph(request.app.state.settings)
    return sync_now(request.app.state.sessions, request.app.state.settings, business_id, user_id)


@router.get("/sync-status")
def status(request: Request, business_id: str, user_id: str = Depends(authorize_business)):
    with request.app.state.sessions() as session:
        business = session.get(Business, business_id)
        source = session.scalar(select(WorkbookSource).where(WorkbookSource.business_id == business_id))
        batches = session.scalars(select(ImportBatch).where(ImportBatch.business_id == business_id)
                                  .order_by(ImportBatch.created_at.desc()).limit(30)).all()
        return {"source_id": source.id if source else None, "kind": source.kind if source else "manual",
                "source_version": business.version, "graph_enabled": request.app.state.settings.graph_enabled,
                "connected": bool(source and source.token_ciphertext and source.linked_user_id == user_id),
                "linked": bool(source and source.item_id and source.linked_user_id == user_id),
                "drive_id": source.drive_id if source and source.linked_user_id == user_id else None,
                "item_id": source.item_id if source and source.linked_user_id == user_id else None,
                "last_import_at": iso_timestamp(source.last_import_at) if source else None,
                "last_checked_at": iso_timestamp(source.last_checked_at) if source else None,
                "max_upload_bytes": request.app.state.settings.max_upload_bytes,
                "imports": [batch_response(batch, business.version) for batch in batches]}


def serialize(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    return value


@router.get("/snapshot")
def snapshot(request: Request, business_id: str, _user_id: str = Depends(authorize_business)):
    # Module integration contract: typed values, dates as ISO strings, exact decimals as strings.
    with request.app.state.sessions.begin() as session:
        if session.bind.dialect.name == "sqlite":
            # sqlite's legacy driver otherwise starts transactions only on writes.
            session.connection().exec_driver_sql("BEGIN")
        business = session.scalar(select(Business).where(Business.id == business_id).with_for_update(read=True))
        batch = session.scalar(select(ImportBatch).where(
            ImportBatch.business_id == business_id, ImportBatch.status == "success"
        ).order_by(ImportBatch.created_at.desc(), ImportBatch.id.desc()).limit(1))
        tables = {}
        for name, table in PROJECTIONS.items():
            tables[name] = [{key: serialize(value) for key, value in row.items() if key not in ("source_id", "batch_id")}
                            for row in session.execute(select(table).where(table.c.business_id == business_id)).mappings()]
        return {"business_id": business_id, "source_version": business.version, "tables": tables,
                "synthetic": is_synthetic_seed_workbook(business_id, batch.content_hash if batch else None)}


@router.get("/template")
def template(request: Request, business_id: str, _user_id: str = Depends(authorize_business)):
    path = ROOT / "shared/workbook-template/shay-sett-template.xlsx"
    if not path.exists():
        raise APIError(503, "TEMPLATE_UNAVAILABLE", "Workbook template has not been generated.")
    workbook = load_workbook(path)
    try:
        for sheet_name in ("Loans", "LoanPayments"):
            if sheet_name in workbook.sheetnames:
                workbook.remove(workbook[sheet_name])
        for sheet in workbook:
            for cell in reversed(list(sheet[1])):
                if cell.value == "business_id":
                    sheet.delete_cols(cell.column)
        if "ReadMe" in workbook.sheetnames:
            notes = workbook["ReadMe"]
            notes["A1"] = "Shay Sett workbook"
            for row in reversed(list(notes.iter_rows())):
                if str(row[0].value).strip().lower() == "onedrive":
                    notes.delete_rows(row[0].row)
            for row in workbook["ReadMe"]:
                for cell in row:
                    if isinstance(cell.value, str) and "business_id" in cell.value:
                        cell.value = (
                            "Enter business_name, currency (MMK), opening_date and opening_cash_mmk."
                            if "opening_cash_mmk" in cell.value else
                            "Uploads are assigned to your selected workspace automatically."
                        )
        for sheet_name, spec in SCHEMA.items():
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            for column in range(sheet.max_column, 0, -1):
                if sheet.cell(1, column).value in spec.get("server_fields", []):
                    sheet.delete_cols(column)
            headers = {cell.value for cell in sheet[1] if cell.value is not None}
            for field in spec.get("optional_fields", []):
                if field in headers:
                    continue
                column = sheet.max_column + 1
                cell = sheet.cell(1, column, field)
                if column > 1:
                    cell._style = copy(sheet.cell(1, column - 1)._style)
                sheet.column_dimensions[cell.column_letter].width = max(
                    sheet.column_dimensions[cell.column_letter].width or 0, 20
                )
        with request.app.state.sessions() as session:
            business = session.get(Business, business_id)
            settings = workbook["Settings"]
            columns = {cell.value: cell.column for cell in settings[1]}
            settings.cell(2, columns["business_name"], business.name)
            settings.cell(2, columns["currency"], business.currency or "MMK")
        buffer = BytesIO()
        workbook.save(buffer)
    finally:
        workbook.close()
    return Response(buffer.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="shay-sett-template.xlsx"'})


def cloud_client(request, business_id, user_id):
    with request.app.state.sessions.begin() as session:
        lock_business(session, business_id)
        source = session.scalar(select(WorkbookSource).where(WorkbookSource.business_id == business_id))
        if not source:
            raise APIError(401, "GRAPH_RECONNECT", "Connect OneDrive first.")
        token = access_token(session, source, user_id, request.app.state.settings)
        return GraphClient(token, request.app.state.settings.max_upload_bytes)


@router.get("/files")
def files(request: Request, business_id: str, folder_id: str | None = None, user_id: str = Depends(authorize_business)):
    client = cloud_client(request, business_id, user_id)
    try:
        return {"items": client.list_files(folder_id)}
    finally:
        client.close()


class Link(BaseModel):
    drive_id: str = Field(min_length=1, max_length=250)
    item_id: str = Field(min_length=1, max_length=250)


@router.post("/link")
def link(body: Link, request: Request, business_id: str, user_id: str = Depends(authorize_business)):
    with request.app.state.sessions.begin() as session:
        lock_business(session, business_id)
        source = session.scalar(select(WorkbookSource).where(WorkbookSource.business_id == business_id))
        if not source:
            raise APIError(401, "GRAPH_RECONNECT", "Connect OneDrive first.")
        token = access_token(session, source, user_id, request.app.state.settings)
        client = GraphClient(token, request.app.state.settings.max_upload_bytes)
        try:
            metadata = client.metadata(body.drive_id, body.item_id)
        finally:
            client.close()
        if "file" not in metadata or not metadata.get("name", "").lower().endswith(".xlsx"):
            raise APIError(400, "FILE_TYPE_INVALID", "Choose an .xlsx workbook.")
        source.drive_id, source.item_id, source.kind, source.last_etag = body.drive_id, body.item_id, "onedrive", None
    return {"status": "linked"}


@router.delete("/link")
def unlink(request: Request, business_id: str, _user_id: str = Depends(authorize_business)):
    with request.app.state.sessions.begin() as session:
        lock_business(session, business_id)
        source = session.scalar(select(WorkbookSource).where(WorkbookSource.business_id == business_id))
        if source:
            source.token_ciphertext = source.drive_id = source.item_id = source.linked_user_id = source.last_etag = None
            source.kind = "manual"
    return {"status": "disconnected", "message": "Saved Microsoft tokens removed. Imported records retained. Revoke app consent in Microsoft to revoke the grant."}
