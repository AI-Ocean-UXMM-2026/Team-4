"""Untrusted XLSX input: bounded ZIP expansion, no formulas, no external links."""
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .schemas import SCHEMA, ParsedWorkbook, WorkbookRow, issue

MAX_ROWS = 10000
MAX_TOTAL_ROWS = 30000
MAX_EXPANDED_BYTES = 40 * 1024 * 1024


def _value(value, kind):
    if value is None or isinstance(value, bool):
        raise ValueError("A value is required")
    if kind == "positive_months":
        result = Decimal(str(value).strip())
        if not result.is_finite() or result != result.to_integral_value() or not 1 <= result <= 1200:
            raise ValueError("Enter a whole number of months between 1 and 1200")
        return int(result)
    if kind == "date":
        if isinstance(value, datetime):
            if value.time().isoformat() != "00:00:00":
                raise ValueError("Use a date without a time")
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return date.fromisoformat(str(value).strip()).isoformat()
    if kind in ("money", "positive_money", "quantity", "positive_quantity"):
        try:
            result = Decimal(str(value).strip())
        except InvalidOperation:
            raise ValueError("Enter a numeric value without currency symbols or separators")
        precision = 2 if "money" in kind else 3
        if not result.is_finite() or result < 0 or result > Decimal("999999999999"):
            raise ValueError("Enter a finite non-negative number up to 999999999999")
        if kind.startswith("positive") and result == 0:
            raise ValueError("Value must be greater than zero")
        if result != result.quantize(Decimal(10) ** -precision):
            raise ValueError(f"Use at most {precision} decimal places")
        return format(result.quantize(Decimal(10) ** -precision), "f")
    text = str(value).strip()
    if not text or len(text) > 500:
        raise ValueError("Enter text between 1 and 500 characters")
    if kind == "id" and (not isinstance(value, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,79}", text)):
        raise ValueError("IDs must be text using letters, numbers, underscores or hyphens (80 characters max)")
    if kind.startswith("enum:") and text not in kind[5:].split(","):
        raise ValueError("Choose one of: " + kind[5:])
    return text


def parse_workbook(uploaded_bytes: bytes, *, business_id: str | None = None) -> ParsedWorkbook:
    parsed = ParsedWorkbook()
    server_fields = {"business_id"} if business_id is not None else set()
    try:
        with ZipFile(BytesIO(uploaded_bytes)) as archive:
            files = archive.infolist()
            if len(files) > 500 or sum(x.file_size for x in files) > MAX_EXPANDED_BYTES:
                raise ValueError("Workbook expands beyond the safe processing limit")
            if any(x.filename.startswith("xl/externalLinks/") or x.filename.endswith("vbaProject.bin") for x in files):
                raise ValueError("External workbook links and macros are not supported")
        workbook = load_workbook(BytesIO(uploaded_bytes), read_only=True, data_only=False, keep_links=False)
        try:
            total = 0
            for name, spec in SCHEMA.items():
                if name not in workbook.sheetnames:
                    if not spec.get("optional"):
                        parsed.errors.append(issue(name, 1, "sheet", "Required sheet is missing"))
                    continue
                sheet = workbook[name]
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise ValueError(f"{name} exceeds the {MAX_ROWS} row limit")
                if sheet.max_column and sheet.max_column > 50:
                    raise ValueError(f"{name} has too many columns")
                rows = sheet.iter_rows()
                header = [str(c.value).strip() if c.value is not None else "" for c in next(rows, [])]
                if len([h for h in header if h]) != len(set(h for h in header if h)):
                    parsed.errors.append(issue(name, 1, "header", "Duplicate column headers"))
                optional_fields = set(spec.get("optional_fields", []))
                missing = set(spec["fields"]) - optional_fields - server_fields - set(header)
                for key in sorted(missing):
                    parsed.errors.append(issue(name, 1, key, "Required column is missing"))
                parsed.sheets[name] = []
                if missing:
                    continue
                for number, cells in enumerate(rows, start=2):
                    if number > MAX_ROWS + 1:
                        raise ValueError(f"{name} exceeds the row limit")
                    if not any(c.value is not None and header[i] not in server_fields
                               for i, c in enumerate(cells)):
                        continue
                    total += 1
                    if total > MAX_TOTAL_ROWS:
                        raise ValueError("Workbook exceeds the total row limit")
                    values = {}
                    for key, kind in spec["fields"].items():
                        if key in server_fields:
                            # Only the authorized request selects the destination workspace.
                            values[key] = business_id
                            continue
                        if key not in header:
                            values[key] = None
                            continue
                        cell = cells[header.index(key)]
                        if key in optional_fields and cell.value is None:
                            values[key] = None
                            continue
                        try:
                            if cell.data_type == "f":
                                raise ValueError("Formula inputs are unsupported; paste the calculated value")
                            value = _value(cell.value, kind)
                            # Workbook sources may call an initially part-paid
                            # sale/purchase "partial". In the ledger this is a
                            # credit transaction with an initial settlement.
                            if key == "payment_type" and value == "partial":
                                value = "credit"
                            if key == "payment_status" and value == "partial":
                                value = "partial"
                            values[key] = value
                        except (ValueError, TypeError, OverflowError) as exc:
                            parsed.errors.append(issue(name, number, key, str(exc)))
                    for key, kind in spec.get("payload_fields", {}).items():
                        if key not in header:
                            continue
                        cell = cells[header.index(key)]
                        if cell.value is None:
                            continue
                        try:
                            if cell.data_type == "f":
                                raise ValueError("Formula inputs are unsupported; paste the calculated value")
                            values[key] = _value(cell.value, kind)
                        except (ValueError, TypeError, OverflowError, InvalidOperation) as exc:
                            parsed.errors.append(issue(name, number, key, str(exc)))
                    parsed.sheets[name].append(WorkbookRow(number, values))
                    if len(parsed.errors) >= 200:
                        raise ValueError("Too many validation errors; correct the reported rows first")
            unknown = set(workbook.sheetnames) - set(SCHEMA) - {"ReadMe"}
            if unknown:
                parsed.warnings.append(issue("Workbook", 0, "sheet", "Ignored extra sheets: " + ", ".join(sorted(unknown))))
        finally:
            workbook.close()
    except (BadZipFile, KeyError, ValueError, OSError, TypeError) as exc:
        parsed.errors.append(issue("Workbook", 0, "file", str(exc) if isinstance(exc, ValueError) else "Unreadable .xlsx workbook"))
    except Exception:
        # XML/package errors are input failures. Never leak workbook contents or parser internals.
        parsed.errors.append(issue("Workbook", 0, "file", "Invalid or unsupported .xlsx workbook"))
    return parsed
